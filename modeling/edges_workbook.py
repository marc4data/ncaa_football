"""One workbook for a week: our numbers, the pack model's, CFBD's and the market's, side by side.

    python -m modeling.weekly --season 2026 --week 5          # first: our predictions for the week
    python -m modeling.edges_workbook --season 2026 --week 5  # then: the workbook

Writes `model_outputs/cfdb_edges_<season>_week<NN>.xlsx` — gitignored, never loaded anywhere, and it
must stay local: the pack model's numbers are built from the licensed pack's features.

SHEETS
    How to read      definitions, sign conventions, and what the rule test found
    Slate            one row per game: the market line and total, each team's market-implied points,
                     ours and the pack model's predicted points, the gap from implied, CFBD's pregame
                     win probability, and flags
    Over board       games where a model puts BOTH teams above their implied points, sorted by the
                     weaker side's gap, with each team's season-to-date scoring against the market
    Team tendencies  per team per season: scoring against implied points, for and against, O/U, ATS
    History          every team-game with a market line, 2013 → last completed week, with any model
                     prediction we hold for it; filter the `team` column to research one team
    Rule test        the pre-registered walk-forward test of the both-over rule (modeling.over_rule)

SIGN CONVENTION (everywhere): spread and margin are AWAY − HOME, CFBD's home line. −7 means the home
team is favoured by 7. Every such number has a text column beside it saying who is favoured.

WHERE EACH NUMBER COMES FROM
    ours            model_outputs/cfdb_wtc_c1_own_<season>_week<NN>.csv  (modeling.weekly)
    pack model      C2 — boosted trees on team points (R-2420), trained here on every pack season before
                    this one and scored on cfdb_model_pack/training_data_<season>_week<NN>.csv
    market          warehouse `marts.fct_betting_line`: each book's latest line, median across books
    CFBD            warehouse `marts.fct_game_pregame_wp`: latest home win probability per game
    history         disk `data/raw/lines` 2013–2025 (disk_source.lines) + warehouse for this season
    history preds   model_outputs/cfdb_over_rule_*.csv (walk-forward) and each weekly CSV
"""
import argparse
import os
from pathlib import Path
from typing import Optional

import numpy as np
import pandas as pd

from modeling import over_rule
from modeling.export import OUTPUT_DIR

PACK_DIR = Path(__file__).resolve().parents[1] / "cfdb_model_pack"
HISTORY_FROM = 2013


# ---------------------------------------------------------------- inputs

def _connect():
    import psycopg2
    return psycopg2.connect(host=os.environ["CFDB_WAREHOUSE_HOST"], port=int(os.environ["CFDB_WAREHOUSE_PORT"]),
                            user=os.environ.get("PG_USER", "cfdb"), password=os.environ.get("PG_PASSWORD", "cfdb"),
                            dbname=os.environ.get("PG_DB", "cfdb"), options="-c default_transaction_read_only=on")


def _read(conn, sql: str, params: Optional[dict] = None) -> pd.DataFrame:
    cur = conn.cursor()
    cur.execute(sql, params or {})
    out = pd.DataFrame(cur.fetchall(), columns=[c[0] for c in cur.description])
    for col in out.columns:          # Postgres NUMERIC arrives as Decimal
        if out[col].dtype == object and len(out[col].dropna()) and hasattr(out[col].dropna().iloc[0], "as_tuple"):
            out[col] = out[col].astype(float)
    return out


def our_week(season: int, week: int) -> pd.DataFrame:
    from modeling.weekly import output_name
    path = OUTPUT_DIR / output_name(season, week)
    if not path.exists():
        raise SystemExit(f"refusing: {path.name} is missing — run `python -m modeling.weekly --season {season} "
                         f"--week {week}` first")
    f = pd.read_csv(path)
    return f.rename(columns={"predicted_home_points": "ours_home", "predicted_away_points": "ours_away"})[
        ["game_id", "start_date", "home_team", "away_team", "home_conference", "away_conference", "neutral_site",
         "ours_home", "ours_away"]]


def pack_week(season: int, week: int) -> Optional[pd.DataFrame]:
    """The pack model (C2) on the week's pack file, trained on every earlier pack season."""
    path = PACK_DIR / f"training_data_{season}_week{week:02d}.csv"
    if not path.exists():
        return None
    from modeling.data import load_pack
    from modeling.models import predict
    from modeling.tune import check_fold

    train = load_pack()
    train = train[(train["season_type"] == "regular") & (train["week"] >= over_rule.FIRST_WEEK)
                  & (train["season"] < season)]
    check_fold(sorted(train["season"].unique()), season, exam_open=True)   # a live week, not the exam
    score = pd.read_csv(path)
    margin, total = predict(train, score, "paired", "wide", "hgb", "team_points")
    return pd.DataFrame({"game_id": score["id"], "pack_home": (total - margin) / 2, "pack_away": (total + margin) / 2})


def market_now(conn, season: int) -> pd.DataFrame:
    from modeling.weekly import market_lines
    return market_lines(conn, [season]).rename(columns={"id": "game_id"})


def cfbd_wp(conn, season: int, week: int) -> pd.DataFrame:
    return _read(conn, """
        select distinct on (w.game_id) w.game_id, w.home_win_probability as cfbd_home_wp
        from marts.fct_game_pregame_wp w join staging.stg_games g on g.game_id = w.game_id
        where g.season = %(s)s and g.week = %(w)s
        order by w.game_id, w.snapshot_ts desc""", {"s": season, "w": week})


def history_games(conn, season: int) -> pd.DataFrame:
    """Every completed regular-season game with a market line: disk to last season, warehouse this one."""
    from modeling import disk_source
    disk = disk_source.lines()
    disk = disk[(disk["season_type"] == "regular") & (disk["season"] >= HISTORY_FROM) & (disk["season"] < season)]
    now = _read(conn, """
        select game_id, season, week, start_date::text as start_date, home_team, away_team, home_conference,
               away_conference, home_classification, away_classification, home_points, away_points
        from staging.stg_games
        where season = %(s)s and season_type = 'regular' and is_completed""", {"s": season})
    now = now.merge(market_now(conn, season)[["game_id", "spread", "market_total"]], on="game_id", how="left")
    games = pd.concat([disk.drop(columns=["season_type", "books"]), now], ignore_index=True)
    return games[games["home_points"].notna() & games["spread"].notna()]


def history_predictions(season: int) -> pd.DataFrame:
    """Every per-game prediction we hold: walk-forward rule-test files, then the live weekly files."""
    parts = []
    for path in sorted(OUTPUT_DIR.glob("cfdb_over_rule_*.csv")):
        f = pd.read_csv(path)
        parts.append(f[["game_id", "model", "pred_home", "pred_away"]])
    for path in sorted(OUTPUT_DIR.glob(f"cfdb_wtc_c1_own_{season}_week*.csv")):
        f = pd.read_csv(path)
        parts.append(pd.DataFrame({"game_id": f["game_id"], "model": "ours",
                                   "pred_home": f["predicted_home_points"], "pred_away": f["predicted_away_points"]}))
    if not parts:
        return pd.DataFrame(columns=["game_id", "ours_home", "ours_away", "pack_home", "pack_away"])
    long = pd.concat(parts, ignore_index=True).drop_duplicates(["game_id", "model"], keep="last")
    wide = long.pivot(index="game_id", columns="model", values=["pred_home", "pred_away"])
    wide.columns = [f"{'ours' if m == 'ours' else 'pack'}_{'home' if p == 'pred_home' else 'away'}"
                    for p, m in wide.columns]
    return wide.reset_index()


# ---------------------------------------------------------------- frames

def _favoured(margin: pd.Series, home: pd.Series, away: pd.Series) -> pd.Series:
    """'Georgia by 7.5' from an away − home number."""
    return pd.Series([("even" if m == 0 else f"{h} by {-m:.1f}" if m < 0 else f"{a} by {m:.1f}") if m == m else ""
                      for m, h, a in zip(margin, home, away)], index=margin.index)


def slate(season: int, week: int, conn) -> pd.DataFrame:
    games = our_week(season, week)
    pack = pack_week(season, week)
    if pack is not None:
        games = games.merge(pack, on="game_id", how="left")
    else:
        games["pack_home"] = games["pack_away"] = np.nan
    games = games.merge(market_now(conn, season)[["game_id", "spread", "market_total", "providers"]],
                        on="game_id", how="left")
    games = games.merge(cfbd_wp(conn, season, week), on="game_id", how="left")

    g = games
    g["implied_home"], g["implied_away"] = over_rule.implied_points(g["spread"], g["market_total"])
    for m in ("ours", "pack"):
        g[f"{m}_gap_away"] = g[f"{m}_away"] - g["implied_away"]
        g[f"{m}_gap_home"] = g[f"{m}_home"] - g["implied_home"]
        g[f"{m}_over_cushion"] = g[[f"{m}_gap_away", f"{m}_gap_home"]].min(axis=1, skipna=False)
        g[f"{m}_margin"] = g[f"{m}_away"] - g[f"{m}_home"]
        g[f"{m}_total"] = g[f"{m}_away"] + g[f"{m}_home"]
    kick = pd.to_datetime(g["start_date"], utc=True).dt.tz_convert("America/Los_Angeles")
    out = pd.DataFrame({
        "kickoff (PT)": kick.dt.strftime("%a %-m/%-d %-I:%M %p"),
        "away": g["away_team"], "home": g["home_team"],
        "neutral": g["neutral_site"].map({True: "yes", False: ""}),
        "away conf": g["away_conference"], "home conf": g["home_conference"],
        "market line": _favoured(g["spread"], g["home_team"], g["away_team"]),
        "market spread (away−home)": g["spread"], "market total": g["market_total"],
        "implied away pts": g["implied_away"], "implied home pts": g["implied_home"],
        "ours away pts": g["ours_away"], "ours home pts": g["ours_home"],
        "ours line": _favoured(g["ours_margin"], g["home_team"], g["away_team"]),
        "ours total": g["ours_total"],
        "ours gap away": g["ours_gap_away"], "ours gap home": g["ours_gap_home"],
        "ours over cushion": g["ours_over_cushion"],
        "pack away pts": g["pack_away"], "pack home pts": g["pack_home"],
        "pack line": _favoured(g["pack_margin"], g["home_team"], g["away_team"]),
        "pack total": g["pack_total"],
        "pack gap away": g["pack_gap_away"], "pack gap home": g["pack_gap_home"],
        "pack over cushion": g["pack_over_cushion"],
        "CFBD home win prob": g["cfbd_home_wp"],
        "BOTH OVER (ours)": np.where(g["ours_over_cushion"] > 0, "yes", ""),
        "BOTH OVER (pack)": np.where(g["pack_over_cushion"] > 0, "yes", ""),
        "BOTH UNDER (ours)": np.where(-g[["ours_gap_away", "ours_gap_home"]].max(axis=1) > 0, "yes", ""),
        "winner differs from market (ours)": np.where(np.sign(g["ours_margin"]) != np.sign(g["spread"]), "yes", ""),
        "books": g["providers"], "game_id": g["game_id"],
    })
    return out.assign(_kick=kick).sort_values(["_kick", "game_id"]).drop(columns="_kick").reset_index(drop=True)


def team_history(games: pd.DataFrame, preds: pd.DataFrame) -> pd.DataFrame:
    """One row per team per game, from that team's side."""
    g = games.merge(preds, on="game_id", how="left")
    g["implied_home"], g["implied_away"] = over_rule.implied_points(g["spread"], g["market_total"])
    rows = []
    for side, other in (("home", "away"), ("away", "home")):
        t = pd.DataFrame({
            "season": g["season"], "week": g["week"], "date": g["start_date"].astype(str).str[:10],
            "team": g[f"{side}_team"], "conference": g[f"{side}_conference"],
            "opponent": g[f"{other}_team"], "site": side,
            "points for": g[f"{side}_points"], "points against": g[f"{other}_points"],
            "implied for": g[f"implied_{side}"], "implied against": g[f"implied_{other}"],
            # the team's own line: home's is the home line as published (spread), away's is its negative
            "team spread": g["spread"] if side == "home" else -g["spread"],
            "market total": g["market_total"],
            "ours pred for": g.get(f"ours_{side}"), "ours pred against": g.get(f"ours_{other}"),
            "pack pred for": g.get(f"pack_{side}"), "pack pred against": g.get(f"pack_{other}"),
            "game_id": g["game_id"],
        })
        rows.append(t)
    h = pd.concat(rows, ignore_index=True)
    h["vs implied for"] = h["points for"] - h["implied for"]
    h["vs implied against"] = h["points against"] - h["implied against"]
    h["margin"] = h["points for"] - h["points against"]
    h["ATS margin"] = h["margin"] + h["team spread"]
    h["covered"] = np.select([h["ATS margin"] > 0, h["ATS margin"] < 0], ["yes", "no"], "push")
    combined = h["points for"] + h["points against"]
    h["combined"] = combined
    over, under = combined > h["market total"], combined < h["market total"]
    h["O/U"] = np.where(h["market total"].isna(), "", np.select([over, under], ["over", "under"], "push"))
    h["total margin"] = combined - h["market total"]
    order = ["season", "week", "date", "team", "conference", "opponent", "site", "points for", "points against",
             "team spread", "ATS margin", "covered", "market total", "combined", "O/U", "total margin",
             "implied for", "vs implied for", "implied against", "vs implied against",
             "ours pred for", "ours pred against", "pack pred for", "pack pred against", "game_id"]
    return h[order].sort_values(["team", "season", "week"]).reset_index(drop=True)


def tendencies(hist: pd.DataFrame) -> pd.DataFrame:
    def one(f):
        ou = f["O/U"]
        return pd.Series({
            "games": len(f),
            "avg vs implied for": f["vs implied for"].mean(),
            "beat implied for": int((f["vs implied for"] > 0).sum()),
            "avg vs implied against": f["vs implied against"].mean(),
            "allowed above implied": int((f["vs implied against"] > 0).sum()),
            "overs": int((ou == "over").sum()), "unders": int((ou == "under").sum()),
            "avg total margin": f["total margin"].mean(),
            "covers": int((f["covered"] == "yes").sum()), "misses": int((f["covered"] == "no").sum()),
            "avg ATS margin": f["ATS margin"].mean(),
        })
    t = hist.groupby(["team", "season"]).apply(one, include_groups=False).reset_index()
    conf = hist.drop_duplicates(["team", "season"], keep="last").set_index(["team", "season"])["conference"]
    t.insert(2, "conference", [conf.get((a, b)) for a, b in zip(t["team"], t["season"])])
    return t.sort_values(["team", "season"]).reset_index(drop=True)


def over_board(sl: pd.DataFrame, tend: pd.DataFrame, season: int) -> pd.DataFrame:
    board = sl[(sl["BOTH OVER (ours)"] == "yes") | (sl["BOTH OVER (pack)"] == "yes")].copy()
    board["weaker-side cushion (best model)"] = board[["ours over cushion", "pack over cushion"]].max(axis=1)
    now = tend[tend["season"] == season].set_index("team")
    for side in ("away", "home"):
        board[f"{side}: avg vs implied for ({season})"] = board[side].map(now["avg vs implied for"])
        board[f"{side}: beat implied for"] = [
            f"{int(now.at[t, 'beat implied for'])} of {int(now.at[t, 'games'])}" if t in now.index else ""
            for t in board[side]]
        board[f"{side}: avg vs implied against ({season})"] = board[side].map(now["avg vs implied against"])
        board[f"{side}: O/U ({season})"] = [
            f"{int(now.at[t, 'overs'])}-{int(now.at[t, 'unders'])}" if t in now.index else "" for t in board[side]]
    cols = ["kickoff (PT)", "away", "home", "market line", "market total", "implied away pts", "implied home pts",
            "ours away pts", "ours home pts", "ours over cushion", "pack away pts", "pack home pts",
            "pack over cushion", "BOTH OVER (ours)", "BOTH OVER (pack)", "weaker-side cushion (best model)"]
    cols += [c for c in board.columns if c.startswith(("away:", "home:"))]
    return board.sort_values("weaker-side cushion (best model)", ascending=False)[cols].reset_index(drop=True)


def rule_test() -> pd.DataFrame:
    parts = []
    for path in sorted(OUTPUT_DIR.glob("cfdb_over_rule_*.csv")):
        name = path.stem.replace("cfdb_over_rule_", "")
        s = over_rule.summarize(pd.read_csv(path))
        s.insert(0, "test", name.replace("_", " "))
        parts.append(s)
    if not parts:
        return pd.DataFrame({"note": ["no rule-test files; run python -m modeling.over_rule"]})
    out = pd.concat(parts, ignore_index=True)
    return out


def how_to_read(season: int, week: int, sl: pd.DataFrame, rt: pd.DataFrame) -> pd.DataFrame:
    primary = rt[(rt.get("direction") == "both over → over") & (rt.get("threshold") == 0)] if "direction" in rt else rt
    lines = [
        ("What this is", f"{season} Week {week}: our model, the pack model (C2), CFBD and the market, side by side."
                         " A research tool — not a pick list."),
        ("Built", f"{pd.Timestamp.now(tz='America/Los_Angeles'):%Y-%m-%d %-I:%M %p} Pacific. The market line is"
                  " the latest at build time; rebuild before kickoff for current lines."),
        ("Sign convention", "Spread and margin are AWAY minus HOME (CFBD's home line): −7 = home favoured by 7."
                            " Every such number has a text column saying who is favoured."),
        ("Implied points", "What the market's line and total imply each team scores: home = (total − spread)/2,"
                           " away = (total + spread)/2."),
        ("Gap", "A model's predicted points minus the implied points, per team. Positive = the model expects that"
                " team to score more than the market does."),
        ("Over cushion", "The SMALLER of the two gaps. Above 0 = the model puts BOTH teams over their implied"
                         " points (Marc's over signal); the weaker side is what decides the over."),
        ("Games", f"{len(sl)} on the slate; ours on {int(sl['ours home pts'].notna().sum())}, pack model on"
                  f" {int(sl['pack home pts'].notna().sum())}, market on {int(sl['market total'].notna().sum())}."),
    ]
    for r in primary.to_dict("records"):
        lines.append((f"Rule test — {r['test']}",
                      f"both over → over fired {r['fires']} times: {r['hits']}-{r['misses']}-{r['pushes']}"
                      f" (over-under-push), {r['hit_rate']:.1%}, 95% range {r['range_low']:.1%}–{r['range_high']:.1%}."
                      f" Break-even 52.4%. Verdict: {r['verdict']}."))
    lines.append(("Reading the rule test", "Pre-registered 2026-09-30 before any result. Across 2018–2024 our model's"
                  " rule ran slightly above break-even but inside the luck band; on 2025, scored once, it did not"
                  " hold. Treat the Over board as a list to research, not a signal that has been proven."))
    lines.append(("History", f"Every team-game with a market line, {HISTORY_FROM} onward. Filter the `team` column."
                  " Model predictions appear where we hold one (walk-forward 2018–2025, live weeks)."))
    return pd.DataFrame(lines, columns=["item", "meaning"])


# ---------------------------------------------------------------- writing

HEADER_FILL, HEADER_FONT = "2F4156", "FFFFFF"
POS_FILL, NEG_FILL, FLAG_FILL = "E3F1E3", "FBE4E4", "FFF3C4"


def _style(ws, frame: pd.DataFrame, freeze: str = "A2", gap_cols=(), flag_cols=(), widths: Optional[dict] = None):
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(bold=True, color=HEADER_FONT)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = freeze
    last = get_column_letter(max(1, len(frame.columns)))
    ws.auto_filter.ref = f"A1:{last}{len(frame) + 1}"
    for i, col in enumerate(frame.columns, 1):
        letter = get_column_letter(i)
        width = (widths or {}).get(col) or min(max(len(str(col)) * 0.55 + 4, 9), 28)
        ws.column_dimensions[letter].width = width
        if pd.api.types.is_float_dtype(frame[col]):
            fmt = "0.0%" if "prob" in col or "rate" in col or col.startswith("range") else "0.0"
            for cell in ws[letter][1:]:
                cell.number_format = fmt
        rng = f"{letter}2:{letter}{len(frame) + 1}"
        if col in gap_cols:
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"],
                                                          fill=PatternFill("solid", fgColor=POS_FILL)))
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                                          fill=PatternFill("solid", fgColor=NEG_FILL)))
        if col in flag_cols:
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"yes"'],
                                                          fill=PatternFill("solid", fgColor=FLAG_FILL)))


def write(path: Path, sheets: dict) -> Path:
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        for name, (frame, opts) in sheets.items():
            frame.to_excel(xl, sheet_name=name, index=False)
            _style(xl.sheets[name], frame, **opts)
    return path


def build(season: int, week: int) -> Path:
    conn = _connect()
    try:
        sl = slate(season, week, conn)
        games = history_games(conn, season)
    finally:
        conn.close()
    hist = team_history(games, history_predictions(season))
    tend = tendencies(hist)
    board = over_board(sl, tend, season)
    rt = rule_test()
    gaps = [c for c in sl.columns if "gap" in c or "cushion" in c]
    flags = [c for c in sl.columns if c.isupper() or c.startswith(("BOTH", "winner"))]
    sheets = {
        "How to read": (how_to_read(season, week, sl, rt), {"widths": {"item": 30, "meaning": 120}}),
        "Slate": (sl, {"freeze": "D2", "gap_cols": gaps, "flag_cols": flags}),
        "Over board": (board, {"freeze": "D2",
                               "gap_cols": [c for c in board.columns if "cushion" in c or "vs implied" in c],
                               "flag_cols": ["BOTH OVER (ours)", "BOTH OVER (pack)"]}),
        "Team tendencies": (tend, {"freeze": "C2", "gap_cols": ["avg vs implied for", "avg vs implied against",
                                                                "avg total margin", "avg ATS margin"]}),
        "History": (hist, {"freeze": "E2", "gap_cols": ["vs implied for", "vs implied against", "ATS margin",
                                                        "total margin"]}),
        "Rule test": (rt, {}),
    }
    return write(OUTPUT_DIR / f"cfdb_edges_{season}_week{week:02d}.xlsx", sheets)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--season", type=int, required=True)
    parser.add_argument("--week", type=int, required=True)
    args = parser.parse_args(argv)
    path = build(args.season, args.week)
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
